# UI/caja_ui.py
from PyQt5.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QLabel, QPushButton,
    QTableWidget, QTableWidgetItem, QGroupBox, QFormLayout,
    QLineEdit, QDoubleSpinBox, QComboBox, QMessageBox,
    QHeaderView, QTabWidget, QGridLayout, QDialog, QSpinBox,
    QDialogButtonBox, QFileDialog
)
from PyQt5.QtCore import Qt, pyqtSignal
from PyQt5.QtGui import QFont, QIntValidator
import sys
import os
from datetime import datetime

sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from database.conexion import DatabaseConnection

class DialogoDenominaciones(QDialog):
    def __init__(self, titulo, detalles, parent=None):
        super().__init__(parent)
        self.setWindowTitle(titulo)
        self.setFixedSize(400, 300)
        layout = QVBoxLayout()
        table = QTableWidget()
        table.setColumnCount(3)
        table.setHorizontalHeaderLabels(["Denominación", "Cantidad", "Subtotal"])
        table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        total = 0
        table.setRowCount(len(detalles))
        for i, (den, cant, subtotal) in enumerate(detalles):
            table.setItem(i, 0, QTableWidgetItem(f"Q {den}.00"))
            table.setItem(i, 1, QTableWidgetItem(str(cant)))
            table.setItem(i, 2, QTableWidgetItem(f"Q {subtotal:.2f}"))
            total += subtotal
        layout.addWidget(table)
        lbl_total = QLabel(f"<b>TOTAL: Q {total:.2f}</b>")
        lbl_total.setAlignment(Qt.AlignRight)
        layout.addWidget(lbl_total)
        btn_box = QDialogButtonBox(QDialogButtonBox.Ok)
        btn_box.accepted.connect(self.accept)
        layout.addWidget(btn_box)
        self.setLayout(layout)

class VentanaCaja(QWidget):
    caja_abierta_signal = pyqtSignal(int)

    def __init__(self, usuario_data):
        super().__init__()
        self.usuario_data = usuario_data
        self.db = DatabaseConnection()
        self.id_caja_actual = None
        self.id_apertura_actual = None
        self.monto_inicial_actual = 0
        self.init_ui()
        self.verificar_estado_caja()
        # Conectar señal de cambio de pestaña para cargar historial automáticamente
        self.tabs.currentChanged.connect(self.on_tab_changed)

    def init_ui(self):
        layout = QVBoxLayout()
        layout.setSpacing(20)

        header = QLabel("Control de Caja")
        header.setFont(QFont("Segoe UI", 18, QFont.Bold))
        layout.addWidget(header)

        self.estado_frame = QLabel()
        self.estado_frame.setStyleSheet("border-radius: 10px; padding: 15px; font-weight: bold;")
        self.estado_frame.setWordWrap(True)
        layout.addWidget(self.estado_frame)

        self.tabs = QTabWidget()
        self.tabs.setStyleSheet("""
            QTabWidget::pane { border: 1px solid #E5E7EB; border-radius: 12px; background-color: white; }
            QTabBar::tab:selected { background-color: #F5C800; border-radius: 8px; padding: 10px 20px; font-weight: bold; }
        """)
        self.tabs.addTab(self.crear_tab_apertura_cierre(), "Control")
        self.tabs.addTab(self.crear_tab_movimientos(), "Movimientos y Resumen")
        self.tabs.addTab(self.crear_tab_historial(), "Historial")

        layout.addWidget(self.tabs)
        self.setLayout(layout)

    def on_tab_changed(self, index):
        if index == 2:  # Índice de la pestaña Historial
            self.cargar_historial()

    def crear_tab_apertura_cierre(self):
        tab = QWidget()
        layout_principal = QHBoxLayout()
        layout_principal.setSpacing(25)
        layout_principal.setContentsMargins(15, 15, 15, 15)

        # Panel izquierdo: conteo de efectivo
        grupo_conteo = QGroupBox("Conteo de efectivo")
        grupo_conteo.setStyleSheet("""
            QGroupBox { font-weight: bold; border: 1px solid #E5E7EB; border-radius: 12px; margin-top: 12px; padding-top: 15px; }
        """)
        ly_conteo = QVBoxLayout()

        grid = QGridLayout()
        self.denominaciones = [200, 100, 50, 20, 10, 5, 1]
        self.inputs_cantidad = {}
        self.labels_subtotal = {}
        validador = QIntValidator(0, 9999)

        for i, den in enumerate(self.denominaciones):
            lbl_den = QLabel(f"Q {den}.00")
            grid.addWidget(lbl_den, i, 0)

            btn_menos = QPushButton("-")
            btn_menos.setFixedSize(30, 30)
            btn_menos.setStyleSheet("background-color: #f3f4f6; border: 1px solid #d1d5db;")
            btn_menos.clicked.connect(lambda checked, d=den: self.ajustar_cantidad(d, -1))
            grid.addWidget(btn_menos, i, 1)

            edit = QLineEdit("0")
            edit.setFixedWidth(60)
            edit.setAlignment(Qt.AlignCenter)
            edit.setValidator(validador)
            edit.textChanged.connect(self.actualizar_totales_desde_conteo)
            self.inputs_cantidad[den] = edit
            grid.addWidget(edit, i, 2)

            btn_mas = QPushButton("+")
            btn_mas.setFixedSize(30, 30)
            btn_mas.setStyleSheet("background-color: #f3f4f6; border: 1px solid #d1d5db;")
            btn_mas.clicked.connect(lambda checked, d=den: self.ajustar_cantidad(d, 1))
            grid.addWidget(btn_mas, i, 3)

            lbl_sub = QLabel("Q 0.00")
            lbl_sub.setAlignment(Qt.AlignRight)
            self.labels_subtotal[den] = lbl_sub
            grid.addWidget(lbl_sub, i, 4)

        ly_conteo.addLayout(grid)

        # --- Sección de monto manual directo ---
        manual_group = QGroupBox("Monto manual (sin billetes)")
        manual_group.setStyleSheet("""
            QGroupBox { font-weight: bold; border: 1px solid #CBD5E1; border-radius: 12px; margin-top: 8px; padding-top: 10px; }
        """)
        manual_layout = QHBoxLayout()
        manual_layout.setSpacing(10)

        self.manual_total = QDoubleSpinBox()
        self.manual_total.setMinimum(0)
        self.manual_total.setMaximum(9999999)
        self.manual_total.setPrefix("Q ")
        self.manual_total.setValue(0)
        self.manual_total.setMinimumWidth(150)
        manual_layout.addWidget(QLabel("Monto único:"))
        manual_layout.addWidget(self.manual_total)

        btn_aplicar_manual = QPushButton("Aplicar monto manual (limpia billetes)")
        btn_aplicar_manual.setStyleSheet("background-color: #3B82F6; color: white; border-radius: 6px; padding: 6px 12px;")
        btn_aplicar_manual.clicked.connect(self.aplicar_monto_manual)
        manual_layout.addWidget(btn_aplicar_manual)

        btn_limpiar_billetes = QPushButton("Limpiar billetes")
        btn_limpiar_billetes.setStyleSheet("background-color: #EF4444; color: white; border-radius: 6px; padding: 6px 12px;")
        btn_limpiar_billetes.clicked.connect(self.limpiar_billetes)
        manual_layout.addWidget(btn_limpiar_billetes)

        manual_group.setLayout(manual_layout)
        ly_conteo.addWidget(manual_group)

        self.lbl_total_conteo = QLabel("<b>TOTAL CONTEO: Q 0.00</b>")
        self.lbl_total_conteo.setAlignment(Qt.AlignRight)
        self.lbl_total_conteo.setStyleSheet("font-size: 14px; margin-top: 10px;")
        ly_conteo.addWidget(self.lbl_total_conteo)

        grupo_conteo.setLayout(ly_conteo)
        layout_principal.addWidget(grupo_conteo, 2)

        # Panel derecho: botones de apertura y cierre
        ly_derecho = QVBoxLayout()
        ly_derecho.setSpacing(15)

        apertura_group = QGroupBox("Apertura de Turno")
        apertura_group.setStyleSheet(grupo_conteo.styleSheet())
        apertura_layout = QFormLayout()
        self.lbl_monto_inicial = QLabel("Q 0.00")
        self.lbl_monto_inicial.setStyleSheet("font-weight: bold; color: #10B981;")
        apertura_layout.addRow("Monto inicial (conteo actual):", self.lbl_monto_inicial)
        self.apertura_btn = QPushButton("Abrir Turno")
        self.apertura_btn.setStyleSheet("background-color: #10B981; color: white; border-radius: 8px; padding: 10px; font-weight: bold;")
        self.apertura_btn.clicked.connect(self.abrir_caja)
        apertura_layout.addRow(self.apertura_btn)
        apertura_group.setLayout(apertura_layout)

        cierre_group = QGroupBox("Cierre de Turno")
        cierre_group.setStyleSheet(grupo_conteo.styleSheet())
        cierre_layout = QFormLayout()
        self.cierre_btn = QPushButton("Cerrar Turno")
        self.cierre_btn.setStyleSheet("background-color: #EF4444; color: white; border-radius: 8px; padding: 10px; font-weight: bold;")
        self.cierre_btn.clicked.connect(self.cerrar_caja)
        cierre_layout.addRow(self.cierre_btn)
        cierre_group.setLayout(cierre_layout)

        ly_derecho.addWidget(apertura_group)
        ly_derecho.addWidget(cierre_group)
        ly_derecho.addStretch()

        layout_principal.addLayout(ly_derecho, 1)
        tab.setLayout(layout_principal)
        return tab

    def ajustar_cantidad(self, den, delta):
        try:
            actual = int(self.inputs_cantidad[den].text() or 0)
            nuevo = max(0, actual + delta)
            self.inputs_cantidad[den].setText(str(nuevo))
        except:
            self.inputs_cantidad[den].setText("0")

    def limpiar_billetes(self):
        for den in self.inputs_cantidad:
            self.inputs_cantidad[den].setText("0")
        self.manual_total.setValue(0)
        self.actualizar_totales_desde_conteo()

    def aplicar_monto_manual(self):
        # Limpiar todos los billetes y usar el monto manual como total
        for den in self.inputs_cantidad:
            self.inputs_cantidad[den].setText("0")
        # El total se actualizará automáticamente porque manual_total.value() se suma en actualizar_totales
        self.actualizar_totales_desde_conteo()
        QMessageBox.information(self, "Monto manual", f"Monto fijado a Q {self.manual_total.value():.2f}")

    def actualizar_totales_desde_conteo(self):
        total = 0
        for den, edit in self.inputs_cantidad.items():
            try:
                cantidad = int(edit.text() or 0)
                subtotal = den * cantidad
                self.labels_subtotal[den].setText(f"Q {subtotal:,.2f}")
                total += subtotal
            except:
                self.labels_subtotal[den].setText("Q 0.00")
        # Sumar monto manual
        total += self.manual_total.value()
        self.lbl_total_conteo.setText(f"<b>TOTAL CONTEO: Q {total:,.2f}</b>")
        self.lbl_monto_inicial.setText(f"Q {total:,.2f}")
        return total

    def obtener_detalles_conteo(self):
        detalles = []
        total = 0
        for den, edit in self.inputs_cantidad.items():
            cant = int(edit.text() or 0)
            if cant > 0:
                subtotal = den * cant
                detalles.append((den, cant, subtotal))
                total += subtotal
        total += self.manual_total.value()
        # Para el cierre, el monto contado será el total (billetes + manual)
        return detalles, total

    def abrir_caja(self):
        if self.id_apertura_actual is not None:
            QMessageBox.warning(self, "Error", "Ya hay un turno abierto. Debe cerrarlo antes de abrir otro.")
            return

        detalles, total = self.obtener_detalles_conteo()
        if total <= 0:
            QMessageBox.warning(self, "Error", "El monto inicial debe ser mayor a cero.")
            return

        fecha_hoy = datetime.now().date()
        caja = self.db.fetch_one("SELECT id_caja FROM caja WHERE fecha = %s", (fecha_hoy,))
        if not caja:
            caja = self.db.fetch_one("INSERT INTO caja (fecha) VALUES (%s) RETURNING id_caja", (fecha_hoy,))
            if not caja:
                QMessageBox.critical(self, "Error", "No se pudo crear la caja del día.")
                return
        id_caja = caja['id_caja']

        apertura_result = self.db.fetch_one("""
            INSERT INTO apertura_cierre 
            (id_caja_fk, id_usuario_fk, fecha_hora_apertura, monto_inicial, estado, observacion_apertura)
            VALUES (%s, %s, NOW(), %s, 'ABIERTO', %s)
            RETURNING id_apertura
        """, (id_caja, self.usuario_data['id_usuario'], total, "Apertura con conteo"))

        if not apertura_result:
            QMessageBox.critical(self, "Error", "No se pudo abrir el turno.")
            return

        id_apertura = apertura_result['id_apertura']

        for den, cant, subtotal in detalles:
            self.db.execute_query("""
                INSERT INTO detalle_apertura (id_apertura_fk, denominacion, cantidad, subtotal)
                VALUES (%s, %s, %s, %s)
            """, (id_apertura, den, cant, subtotal))

        QMessageBox.information(self, "Éxito", f"Turno abierto correctamente con Q {total:,.2f}")
        self.verificar_estado_caja()

    def cerrar_caja(self):
        if not self.id_apertura_actual:
            QMessageBox.warning(self, "Error", "No hay un turno abierto para cerrar.")
            return

        detalles_cierre, monto_contado = self.obtener_detalles_conteo()

        apertura = self.db.fetch_one(
            "SELECT id_caja_fk, fecha_hora_apertura FROM apertura_cierre WHERE id_apertura = %s",
            (self.id_apertura_actual,)
        )
        if not apertura:
            return
        id_caja = apertura['id_caja_fk']
        desde = apertura['fecha_hora_apertura']

        # Ventas normales
        query_normales = """
            SELECT 
                COALESCE(SUM(v.total) FILTER (WHERE v.forma_pago = 'EF' AND v.producto_pagado = TRUE), 0) AS efectivo
            FROM venta v
            JOIN movimiento_caja mc ON v.id_movimiento_fk = mc.id_movimiento
            WHERE mc.id_caja_fk = %s AND mc.fecha_hora >= %s
              AND v.forma_pago != 'MIXTO'
        """
        normales = self.db.fetch_one(query_normales, (id_caja, desde)) or {}
        efectivo_norm = float(normales.get('efectivo', 0))

        query_mixtos = """
            SELECT COALESCE(SUM(dpm.monto) FILTER (WHERE dpm.forma_pago = 'EFECTIVO'), 0) AS efectivo
            FROM detalle_pago_mixto dpm
            JOIN venta v ON dpm.id_venta_fk = v.id_venta
            JOIN movimiento_caja mc ON v.id_movimiento_fk = mc.id_movimiento
            WHERE mc.id_caja_fk = %s AND mc.fecha_hora >= %s
              AND v.producto_pagado = TRUE
        """
        mixtos = self.db.fetch_one(query_mixtos, (id_caja, desde)) or {}
        efectivo_mix = float(mixtos.get('efectivo', 0))

        query_aptos_ef = """
            SELECT COALESCE(SUM(da.monto) FILTER (WHERE da.forma_pago = 'EF'), 0) AS efectivo
            FROM detalle_apartado da
            JOIN movimiento_caja mc ON da.id_movimiento_fk = mc.id_movimiento
            WHERE mc.id_caja_fk = %s AND mc.fecha_hora >= %s
        """
        aptos_ef = self.db.fetch_one(query_aptos_ef, (id_caja, desde)) or {}
        efectivo_apartados = float(aptos_ef.get('efectivo', 0))

        query_otros = """
            SELECT 
                COALESCE(SUM(mc.monto) FILTER (WHERE mc.tipo_movimiento = 'EGRESO'), 0) AS egresos,
                COALESCE(SUM(mc.monto) FILTER (WHERE mc.tipo_movimiento = 'INGRESO'
                    AND v.id_venta IS NULL
                    AND da.id_detalle IS NULL
                    AND mc.descripcion NOT ILIKE '%%MIXTA%%'), 0) AS otros_ingresos
            FROM movimiento_caja mc
            LEFT JOIN venta v ON mc.id_movimiento = v.id_movimiento_fk
            LEFT JOIN detalle_apartado da ON mc.id_movimiento = da.id_movimiento_fk
            WHERE mc.id_caja_fk = %s AND mc.fecha_hora >= %s
        """
        otros = self.db.fetch_one(query_otros, (id_caja, desde)) or {}
        egresos = float(otros.get('egresos', 0))
        otros_ingresos = float(otros.get('otros_ingresos', 0))

        ventas_efectivo = efectivo_norm + efectivo_mix + efectivo_apartados
        monto_esperado = self.monto_inicial_actual + ventas_efectivo + otros_ingresos - egresos
        diferencia = monto_contado - monto_esperado

        resumen = f"""
        <b>RESUMEN DE CIERRE</b><br><br>
        Monto inicial: Q {self.monto_inicial_actual:,.2f}<br>
        Ventas en efectivo: Q {ventas_efectivo:,.2f}<br>
        Otros ingresos: Q {otros_ingresos:,.2f}<br>
        Egresos: Q {egresos:,.2f}<br>
        <b>Efectivo esperado:</b> Q {monto_esperado:,.2f}<br>
        <b>Efectivo contado:</b> Q {monto_contado:,.2f}<br>
        <b>Diferencia:</b> Q {diferencia:,.2f}<br>
        """
        if abs(diferencia) < 0.01:
            resumen += "<span style='color:green'>✓ CAJA CUADRADA</span>"
        elif diferencia > 0:
            resumen += "<span style='color:orange'>⚠️ Sobrante</span>"
        else:
            resumen += "<span style='color:red'>❌ Faltante</span>"

        reply = QMessageBox.question(self, "Confirmar cierre", resumen, QMessageBox.Yes | QMessageBox.No)
        if reply != QMessageBox.Yes:
            return

        query_update = """
            UPDATE apertura_cierre
            SET 
                fecha_hora_cierre = NOW(),
                monto_final = %s,
                monto_esperado = %s,
                diferencia = %s,
                observacion_cierre = %s,
                estado = 'CERRADO',
                id_usuario_cierre_fk = %s
            WHERE id_apertura = %s AND estado = 'ABIERTO'
        """
        obs = f"Cierre. Contado: Q{monto_contado:.2f}, Esperado: Q{monto_esperado:.2f}, Diferencia: Q{diferencia:.2f}"
        exito = self.db.execute_query(
            query_update,
            (monto_contado, monto_esperado, diferencia, obs, self.usuario_data['id_usuario'], self.id_apertura_actual)
        )

        if not exito:
            QMessageBox.critical(self, "Error", "No se pudo cerrar el turno.")
            return

        for den, cant, subtotal in detalles_cierre:
            self.db.execute_query("""
                INSERT INTO detalle_cierre (id_apertura_fk, denominacion, cantidad, subtotal)
                VALUES (%s, %s, %s, %s)
            """, (self.id_apertura_actual, den, cant, subtotal))

        QMessageBox.information(self, "Éxito", "Turno cerrado correctamente")
        self.verificar_estado_caja()

    # ------------------- MÉTODOS DE MOVIMIENTOS Y RESUMEN -------------------
    def crear_tab_movimientos(self):
        tab = QWidget()
        layout = QVBoxLayout()
        layout.setSpacing(15)
        layout.setContentsMargins(10, 10, 10, 10)

        # Resumen del turno en grid de 2 columnas
        resumen_group = QGroupBox("Resumen del Turno Actual")
        resumen_group.setStyleSheet("""
            QGroupBox {
                font-weight: bold;
                border: 1px solid #E5E7EB;
                border-radius: 12px;
                margin-top: 12px;
                padding-top: 15px;
                background-color: #F8FAFC;
            }
        """)
        resumen_layout = QGridLayout()
        resumen_layout.setHorizontalSpacing(20)
        resumen_layout.setVerticalSpacing(8)

        # Primera columna
        resumen_layout.addWidget(QLabel("TOTAL VENTAS DEL TURNO:"), 0, 0)
        self.lbl_total_ventas = QLabel("Q 0.00")
        self.lbl_total_ventas.setStyleSheet("font-size: 16px; font-weight: bold; color: #111827;")
        resumen_layout.addWidget(self.lbl_total_ventas, 0, 1)

        resumen_layout.addWidget(QLabel("💵 Efectivo:"), 1, 0)
        self.lbl_ventas_efectivo = QLabel("Q 0.00")
        self.lbl_ventas_efectivo.setStyleSheet("color: #10B981; font-weight: bold;")
        resumen_layout.addWidget(self.lbl_ventas_efectivo, 1, 1)

        resumen_layout.addWidget(QLabel("💳 Tarjeta:"), 2, 0)
        self.lbl_ventas_tarjeta = QLabel("Q 0.00")
        self.lbl_ventas_tarjeta.setStyleSheet("color: #3B82F6;")
        resumen_layout.addWidget(self.lbl_ventas_tarjeta, 2, 1)

        resumen_layout.addWidget(QLabel("🏦 Transferencia:"), 3, 0)
        self.lbl_ventas_transferencia = QLabel("Q 0.00")
        self.lbl_ventas_transferencia.setStyleSheet("color: #8B5CF6;")
        resumen_layout.addWidget(self.lbl_ventas_transferencia, 3, 1)

        # Segunda columna
        resumen_layout.addWidget(QLabel("📥 Depósito:"), 0, 2)
        self.lbl_ventas_deposito = QLabel("Q 0.00")
        self.lbl_ventas_deposito.setStyleSheet("color: #F59E0B;")
        resumen_layout.addWidget(self.lbl_ventas_deposito, 0, 3)

        resumen_layout.addWidget(QLabel("📦 Cuentas por cobrar:"), 1, 2)
        self.lbl_cuentas_cobrar = QLabel("Q 0.00")
        self.lbl_cuentas_cobrar.setStyleSheet("color: #F59E0B;")
        resumen_layout.addWidget(self.lbl_cuentas_cobrar, 1, 3)

        resumen_layout.addWidget(QLabel("💰 EFECTIVO ACTUAL EN CAJA:"), 2, 2)
        self.lbl_efectivo_actual = QLabel("Q 0.00")
        self.lbl_efectivo_actual.setStyleSheet("font-size: 18px; font-weight: bold; color: #059669;")
        resumen_layout.addWidget(self.lbl_efectivo_actual, 2, 3)

        resumen_layout.addWidget(QLabel("💸 Egresos (gastos/retiros):"), 3, 2)
        self.lbl_egresos = QLabel("Q 0.00")
        self.lbl_egresos.setStyleSheet("color: #EF4444;")
        resumen_layout.addWidget(self.lbl_egresos, 3, 3)

        resumen_group.setLayout(resumen_layout)
        layout.addWidget(resumen_group)

        # Formulario de movimientos en horizontal
        form_group = QGroupBox("Registrar Movimiento Manual")
        form_group.setStyleSheet("""
            QGroupBox {
                font-weight: bold;
                border: 1px solid #E5E7EB;
                border-radius: 12px;
                margin-top: 12px;
                padding-top: 10px;
            }
        """)
        form_layout = QHBoxLayout()
        form_layout.setSpacing(10)

        self.tipo_movimiento = QComboBox()
        self.tipo_movimiento.addItems(["INGRESO", "EGRESO"])
        self.tipo_movimiento.currentTextChanged.connect(self.toggle_gasto_field)
        self.tipo_movimiento.setMinimumWidth(120)
        form_layout.addWidget(QLabel("Tipo:"))
        form_layout.addWidget(self.tipo_movimiento)

        self.tipo_gasto = QComboBox()
        self.tipo_gasto.addItems(["PROVEEDOR", "SUELDOS", "SERVICIOS", "INSUMOS", "DEVOLUCION", "OTRO"])
        self.tipo_gasto.setEnabled(False)
        self.tipo_gasto.setMinimumWidth(120)
        form_layout.addWidget(QLabel("Tipo Gasto:"))
        form_layout.addWidget(self.tipo_gasto)

        self.descripcion_mov = QLineEdit()
        self.descripcion_mov.setPlaceholderText("Descripción")
        self.descripcion_mov.setMinimumWidth(200)
        form_layout.addWidget(QLabel("Descripción:"))
        form_layout.addWidget(self.descripcion_mov)

        self.monto_mov = QDoubleSpinBox()
        self.monto_mov.setMinimum(0)
        self.monto_mov.setMaximum(100000)
        self.monto_mov.setPrefix("Q ")
        self.monto_mov.setMinimumWidth(120)
        form_layout.addWidget(QLabel("Monto:"))
        form_layout.addWidget(self.monto_mov)

        registrar_btn = QPushButton("Registrar Movimiento")
        registrar_btn.setStyleSheet("background-color: #F5C800; border-radius: 8px; padding: 8px 16px; font-weight: bold;")
        registrar_btn.clicked.connect(self.registrar_movimiento)
        form_layout.addWidget(registrar_btn)

        form_group.setLayout(form_layout)
        layout.addWidget(form_group)

        # Tabla de movimientos
        self.movimientos_table = QTableWidget()
        self.movimientos_table.setColumnCount(6)
        self.movimientos_table.setHorizontalHeaderLabels(["Fecha", "Tipo", "Forma Pago", "Descripción", "Monto", "Usuario"])
        self.movimientos_table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        layout.addWidget(self.movimientos_table)

        # Botón exportar
        btn_exportar_turno = QPushButton("Exportar turno actual a Excel")
        btn_exportar_turno.clicked.connect(self.exportar_turno)
        btn_exportar_turno.setStyleSheet("background-color: #F5C800; border-radius: 8px; padding: 6px;")
        layout.addWidget(btn_exportar_turno)

        tab.setLayout(layout)
        return tab

    def toggle_gasto_field(self, tipo):
        self.tipo_gasto.setEnabled(tipo == "EGRESO")

    def registrar_movimiento(self):
        if not self.id_apertura_actual:
            QMessageBox.warning(self, "Error", "Debe abrir un turno primero")
            return

        tipo = self.tipo_movimiento.currentText()
        tipo_gasto = self.tipo_gasto.currentText()
        descripcion = self.descripcion_mov.text().strip()
        monto = self.monto_mov.value()

        if not descripcion or monto <= 0:
            QMessageBox.warning(self, "Error", "Complete todos los campos correctamente")
            return

        query_mov = """
            INSERT INTO movimiento_caja (id_caja_fk, tipo_movimiento, descripcion, monto, fecha_hora, id_usuario_fk)
            VALUES (%s, %s, %s, %s, NOW(), %s) RETURNING id_movimiento
        """
        movimiento = self.db.fetch_one(query_mov,
                                       (self.id_caja_actual, tipo, descripcion, monto, self.usuario_data['id_usuario']))
        if not movimiento:
            QMessageBox.critical(self, "Error", "No se pudo registrar el movimiento")
            return

        if tipo == "EGRESO":
            query_gasto = """
                INSERT INTO gasto (id_movimiento_fk, tipo_gasto, descripcion, monto)
                VALUES (%s, %s, %s, %s)
            """
            if not self.db.execute_query(query_gasto, (movimiento['id_movimiento'], tipo_gasto, descripcion, monto)):
                QMessageBox.critical(self, "Error", "No se pudo registrar el gasto")
                return

        QMessageBox.information(self, "Éxito", "Movimiento registrado")
        self.descripcion_mov.clear()
        self.monto_mov.setValue(0)
        self.cargar_movimientos()
        self.actualizar_resumen_turno()

    # ------------------- MÉTODOS DE HISTORIAL -------------------
    def crear_tab_historial(self):
        tab = QWidget()
        layout = QVBoxLayout()
        self.historial_table = QTableWidget()
        self.historial_table.setColumnCount(8)
        self.historial_table.setHorizontalHeaderLabels([
            "Fecha Apertura", "Fecha Cierre", "Usuario Apertura", "Usuario Cierre",
            "Monto Inicial", "Monto Esperado", "Monto Final", "Diferencia"
        ])
        self.historial_table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        layout.addWidget(self.historial_table)
        # Ya no hay botón actualizar, se carga automáticamente al cambiar de pestaña
        tab.setLayout(layout)
        return tab

    # Los demás métodos (verificar_estado_caja, cargar_movimientos, actualizar_resumen_turno, cargar_historial, exportar_turno, etc.)
    # se mantienen igual que en tu versión anterior. Solo asegúrate de que existan.
    # Aquí incluyo los que faltan por completitud (ya los tenías, los copio desde la vers anterior)

    def verificar_estado_caja(self):
        query = """
            SELECT ac.id_apertura, ac.id_caja_fk, ac.monto_inicial, ac.fecha_hora_apertura,
                   u.nombre as usuario_nombre
            FROM apertura_cierre ac
            JOIN usuario u ON ac.id_usuario_fk = u.id_usuario
            WHERE ac.estado = 'ABIERTO'
            ORDER BY ac.fecha_hora_apertura DESC LIMIT 1
        """
        resultado = self.db.fetch_one(query)
        if resultado:
            self.id_apertura_actual = resultado['id_apertura']
            self.id_caja_actual = resultado['id_caja_fk']
            self.monto_inicial_actual = float(resultado['monto_inicial'])
            self.caja_abierta_signal.emit(self.id_caja_actual)

            detalles_apertura = self.db.fetch_all("""
                SELECT denominacion, cantidad FROM detalle_apertura WHERE id_apertura_fk = %s
            """, (self.id_apertura_actual,))
            for det in detalles_apertura:
                den = det['denominacion']
                if den in self.inputs_cantidad:
                    self.inputs_cantidad[den].setText(str(det['cantidad']))
            self.actualizar_totales_desde_conteo()

            self.estado_frame.setText(
                f" TURNO ABIERTO\n"
                f"Usuario: {resultado['usuario_nombre']}\n"
                f"Monto inicial: Q {self.monto_inicial_actual:,.2f}\n"
                f"Apertura: {resultado['fecha_hora_apertura']}"
            )
            self.estado_frame.setStyleSheet("background-color: #D1FAE5; color: #059669; border-radius: 10px; padding: 15px;")
            self.apertura_btn.setEnabled(False)
            self.cierre_btn.setEnabled(True)
            self.cargar_movimientos()
            self.actualizar_resumen_turno()
        else:
            self.id_apertura_actual = None
            self.id_caja_actual = None
            self.monto_inicial_actual = 0
            for den, edit in self.inputs_cantidad.items():
                edit.setText("0")
            self.manual_total.setValue(0)
            self.actualizar_totales_desde_conteo()

            ultimo_cierre = self.db.fetch_one("""
                SELECT ac.monto_final, ac.fecha_hora_cierre, u.nombre as usuario_nombre
                FROM apertura_cierre ac
                JOIN usuario u ON ac.id_usuario_fk = u.id_usuario
                WHERE ac.estado = 'CERRADO'
                ORDER BY ac.fecha_hora_cierre DESC LIMIT 1
            """)
            if ultimo_cierre:
                monto_cierre = float(ultimo_cierre['monto_final'])
                fecha = ultimo_cierre['fecha_hora_cierre']
                usuario = ultimo_cierre['usuario_nombre']
                self.estado_frame.setText(
                    f" TURNO CERRADO\n"
                    f"Último cierre: Q {monto_cierre:,.2f}\n"
                    f"Fecha: {fecha}\n"
                    f"Cajero: {usuario}"
                )
                if not hasattr(self, 'btn_den_cierre'):
                    self.btn_den_cierre = QPushButton("Ver denominaciones del último cierre")
                    self.btn_den_cierre.clicked.connect(self.ver_denominaciones_ultimo_cierre)
                    self.layout().insertWidget(2, self.btn_den_cierre)
                else:
                    self.btn_den_cierre.show()
            else:
                self.estado_frame.setText(" TURNO CERRADO\nNo hay cierres previos")
                if hasattr(self, 'btn_den_cierre'):
                    self.btn_den_cierre.hide()

            self.estado_frame.setStyleSheet("background-color: #FEE2E2; color: #DC2626; border-radius: 10px; padding: 15px;")
            self.apertura_btn.setEnabled(True)
            self.cierre_btn.setEnabled(False)
            # No llamamos a cargar_historial aquí porque se cargará al cambiar de pestaña
            self.movimientos_table.setRowCount(0)
            # Limpiar resumen
            self.lbl_total_ventas.setText("Q 0.00")
            self.lbl_ventas_efectivo.setText("Q 0.00")
            self.lbl_ventas_tarjeta.setText("Q 0.00")
            self.lbl_ventas_transferencia.setText("Q 0.00")
            self.lbl_ventas_deposito.setText("Q 0.00")
            self.lbl_cuentas_cobrar.setText("Q 0.00")
            self.lbl_efectivo_actual.setText("Q 0.00")
            self.lbl_egresos.setText("Q 0.00")

    def ver_denominaciones_ultimo_cierre(self):
        query = """
            SELECT dc.denominacion, dc.cantidad, dc.subtotal
            FROM detalle_cierre dc
            JOIN apertura_cierre ac ON dc.id_apertura_fk = ac.id_apertura
            WHERE ac.estado = 'CERRADO'
            ORDER BY ac.fecha_hora_cierre DESC
            LIMIT 100
        """
        detalles = self.db.fetch_all(query)
        if not detalles:
            QMessageBox.information(self, "Información", "No hay detalles de denominaciones para el último cierre")
            return
        detalles_list = [(d['denominacion'], d['cantidad'], float(d['subtotal'])) for d in detalles]
        dialog = DialogoDenominaciones("Denominaciones del último cierre", detalles_list, self)
        dialog.exec_()

    def cargar_movimientos(self):
        if not self.id_apertura_actual:
            self.movimientos_table.setRowCount(0)
            return

        query = """
            SELECT 
                mc.fecha_hora, 
                mc.tipo_movimiento, 
                mc.descripcion, 
                mc.monto, 
                u.nombre,
                CASE 
                    WHEN v.id_venta IS NOT NULL AND v.forma_pago != 'MIXTO' THEN v.forma_pago
                    WHEN v.id_venta IS NOT NULL AND v.forma_pago = 'MIXTO' THEN (
                        SELECT forma_pago FROM detalle_pago_mixto dpm 
                        WHERE dpm.id_venta_fk = v.id_venta 
                          AND dpm.monto = mc.monto
                        LIMIT 1
                    )
                    WHEN da.id_detalle IS NOT NULL THEN da.forma_pago
                    ELSE NULL
                END AS forma_pago_detalle
            FROM movimiento_caja mc
            JOIN apertura_cierre ac ON mc.id_caja_fk = ac.id_caja_fk
            LEFT JOIN usuario u ON mc.id_usuario_fk = u.id_usuario
            LEFT JOIN venta v ON mc.id_movimiento = v.id_movimiento_fk
            LEFT JOIN detalle_apartado da ON mc.id_movimiento = da.id_movimiento_fk
            WHERE ac.id_apertura = %s
              AND mc.fecha_hora >= ac.fecha_hora_apertura
              AND (ac.fecha_hora_cierre IS NULL OR mc.fecha_hora <= ac.fecha_hora_cierre)
            ORDER BY mc.fecha_hora DESC
        """
        movs = self.db.fetch_all(query, (self.id_apertura_actual,))
        self.movimientos_table.setRowCount(len(movs))
        for i, m in enumerate(movs):
            self.movimientos_table.setItem(i, 0, QTableWidgetItem(str(m['fecha_hora'])[:19]))
            self.movimientos_table.setItem(i, 1, QTableWidgetItem(m['tipo_movimiento']))
            fp = m.get('forma_pago_detalle')
            if fp:
                forma_texto = {
                    'EF': '💵 Efectivo',
                    'TC/TD': '💳 Tarjeta',
                    'TF': '🏦 Transferencia',
                    'DP': '📥 Depósito',
                    'EFECTIVO': '💵 Efectivo',
                    'TARJETA': '💳 Tarjeta',
                    'TRANSFERENCIA': '🏦 Transferencia',
                    'DEPOSITO': '📥 Depósito',
                }.get(fp, fp)
            else:
                forma_texto = '—'
            self.movimientos_table.setItem(i, 2, QTableWidgetItem(forma_texto))
            self.movimientos_table.setItem(i, 3, QTableWidgetItem(m['descripcion']))
            self.movimientos_table.setItem(i, 4, QTableWidgetItem(f"Q {float(m['monto']):,.2f}"))
            self.movimientos_table.setItem(i, 5, QTableWidgetItem(m['nombre']))
        self.actualizar_resumen_turno()

    def actualizar_resumen_turno(self):
        if not self.id_apertura_actual:
            return

        apertura = self.db.fetch_one(
            "SELECT id_caja_fk, fecha_hora_apertura FROM apertura_cierre WHERE id_apertura = %s",
            (self.id_apertura_actual,)
        )
        if not apertura:
            return
        id_caja = apertura['id_caja_fk']
        desde = apertura['fecha_hora_apertura']

        query_normales = """
            SELECT 
                COALESCE(SUM(v.total) FILTER (WHERE v.forma_pago = 'EF' AND v.producto_pagado = TRUE), 0) AS efectivo,
                COALESCE(SUM(v.total) FILTER (WHERE v.forma_pago = 'TC/TD' AND v.producto_pagado = TRUE), 0) AS tarjeta,
                COALESCE(SUM(v.total) FILTER (WHERE v.forma_pago = 'TF' AND v.producto_pagado = TRUE), 0) AS transferencia,
                COALESCE(SUM(v.total) FILTER (WHERE v.forma_pago = 'DP' AND v.producto_pagado = TRUE), 0) AS deposito,
                COALESCE(SUM(v.total) FILTER (WHERE v.producto_pagado = FALSE), 0) AS cuentas_cobrar
            FROM venta v
            JOIN movimiento_caja mc ON v.id_movimiento_fk = mc.id_movimiento
            WHERE mc.id_caja_fk = %s AND mc.fecha_hora >= %s
              AND v.forma_pago != 'MIXTO'
        """
        normales = self.db.fetch_one(query_normales, (id_caja, desde)) or {}
        efectivo = float(normales.get('efectivo', 0))
        tarjeta = float(normales.get('tarjeta', 0))
        transferencia = float(normales.get('transferencia', 0))
        deposito = float(normales.get('deposito', 0))
        cuentas_cobrar = float(normales.get('cuentas_cobrar', 0))

        query_mixtos = """
            SELECT 
                COALESCE(SUM(dpm.monto) FILTER (WHERE dpm.forma_pago = 'EFECTIVO'), 0) AS efectivo,
                COALESCE(SUM(dpm.monto) FILTER (WHERE dpm.forma_pago = 'TARJETA'), 0) AS tarjeta,
                COALESCE(SUM(dpm.monto) FILTER (WHERE dpm.forma_pago = 'TRANSFERENCIA'), 0) AS transferencia,
                COALESCE(SUM(dpm.monto) FILTER (WHERE dpm.forma_pago = 'DEPOSITO'), 0) AS deposito
            FROM detalle_pago_mixto dpm
            JOIN venta v ON dpm.id_venta_fk = v.id_venta
            JOIN movimiento_caja mc ON v.id_movimiento_fk = mc.id_movimiento
            WHERE mc.id_caja_fk = %s AND mc.fecha_hora >= %s
              AND v.producto_pagado = TRUE
        """
        mixtos = self.db.fetch_one(query_mixtos, (id_caja, desde)) or {}
        efectivo += float(mixtos.get('efectivo', 0))
        tarjeta += float(mixtos.get('tarjeta', 0))
        transferencia += float(mixtos.get('transferencia', 0))
        deposito += float(mixtos.get('deposito', 0))

        query_apartados = """
            SELECT 
                COALESCE(SUM(da.monto) FILTER (WHERE da.forma_pago = 'EF'), 0) AS efectivo,
                COALESCE(SUM(da.monto) FILTER (WHERE da.forma_pago = 'TC/TD'), 0) AS tarjeta,
                COALESCE(SUM(da.monto) FILTER (WHERE da.forma_pago = 'TF'), 0) AS transferencia,
                COALESCE(SUM(da.monto) FILTER (WHERE da.forma_pago = 'DP'), 0) AS deposito
            FROM detalle_apartado da
            JOIN movimiento_caja mc ON da.id_movimiento_fk = mc.id_movimiento
            WHERE mc.id_caja_fk = %s AND mc.fecha_hora >= %s
        """
        aptos = self.db.fetch_one(query_apartados, (id_caja, desde)) or {}
        efectivo += float(aptos.get('efectivo', 0))
        tarjeta += float(aptos.get('tarjeta', 0))
        transferencia += float(aptos.get('transferencia', 0))
        deposito += float(aptos.get('deposito', 0))

        query_otros = """
            SELECT 
                COALESCE(SUM(mc.monto) FILTER (WHERE mc.tipo_movimiento = 'EGRESO'), 0) AS egresos,
                COALESCE(SUM(mc.monto) FILTER (WHERE mc.tipo_movimiento = 'INGRESO'
                    AND v.id_venta IS NULL
                    AND da.id_detalle IS NULL
                    AND mc.descripcion NOT ILIKE '%%MIXTA%%'), 0) AS otros_ingresos
            FROM movimiento_caja mc
            LEFT JOIN venta v ON mc.id_movimiento = v.id_movimiento_fk
            LEFT JOIN detalle_apartado da ON mc.id_movimiento = da.id_movimiento_fk
            WHERE mc.id_caja_fk = %s AND mc.fecha_hora >= %s
        """
        otros = self.db.fetch_one(query_otros, (id_caja, desde)) or {}
        egresos = float(otros.get('egresos', 0))
        otros_ingresos = float(otros.get('otros_ingresos', 0))

        total_ventas_pagadas = efectivo + tarjeta + transferencia + deposito
        total_ventas_turno = total_ventas_pagadas + cuentas_cobrar
        efectivo_en_caja = self.monto_inicial_actual + efectivo + otros_ingresos - egresos

        self.lbl_total_ventas.setText(f"Q {total_ventas_turno:,.2f}")
        self.lbl_ventas_efectivo.setText(f"Q {efectivo:,.2f}")
        self.lbl_ventas_tarjeta.setText(f"Q {tarjeta:,.2f}")
        self.lbl_ventas_transferencia.setText(f"Q {transferencia:,.2f}")
        self.lbl_ventas_deposito.setText(f"Q {deposito:,.2f}")
        self.lbl_cuentas_cobrar.setText(f"Q {cuentas_cobrar:,.2f}")
        self.lbl_efectivo_actual.setText(f"Q {efectivo_en_caja:,.2f}")
        self.lbl_egresos.setText(f"Q {egresos:,.2f}")

    def cargar_historial(self):
        query = """
            SELECT ac.fecha_hora_apertura, ac.fecha_hora_cierre, 
                   ua.nombre as usuario_apertura, uc.nombre as usuario_cierre,
                   ac.monto_inicial, ac.monto_esperado, ac.monto_final, ac.diferencia
            FROM apertura_cierre ac
            JOIN usuario ua ON ac.id_usuario_fk = ua.id_usuario
            LEFT JOIN usuario uc ON ac.id_usuario_cierre_fk = uc.id_usuario
            WHERE ac.estado = 'CERRADO'
            ORDER BY ac.fecha_hora_apertura DESC
        """
        hist = self.db.fetch_all(query)
        self.historial_table.setRowCount(len(hist))
        for i, h in enumerate(hist):
            self.historial_table.setItem(i, 0, QTableWidgetItem(str(h['fecha_hora_apertura'])[:19]))
            self.historial_table.setItem(i, 1, QTableWidgetItem(str(h['fecha_hora_cierre'])[:19]) if h['fecha_hora_cierre'] else "-")
            self.historial_table.setItem(i, 2, QTableWidgetItem(h['usuario_apertura']))
            self.historial_table.setItem(i, 3, QTableWidgetItem(h['usuario_cierre'] if h['usuario_cierre'] else "-"))
            self.historial_table.setItem(i, 4, QTableWidgetItem(f"Q {float(h['monto_inicial']):,.2f}"))
            esperado = f"Q {float(h['monto_esperado']):,.2f}" if h['monto_esperado'] is not None else "-"
            self.historial_table.setItem(i, 5, QTableWidgetItem(esperado))
            final = f"Q {float(h['monto_final']):,.2f}" if h['monto_final'] else "-"
            self.historial_table.setItem(i, 6, QTableWidgetItem(final))
            diff = f"Q {float(h['diferencia']):,.2f}" if h['diferencia'] is not None else "-"
            self.historial_table.setItem(i, 7, QTableWidgetItem(diff))

    def exportar_turno(self):
        if not self.id_apertura_actual:
            QMessageBox.warning(self, "Sin turno", "No hay un turno abierto.")
            return
        from services.reporte_service import ReporteService
        rs = ReporteService()
        movimientos = rs.obtener_movimientos_turno(self.id_apertura_actual)
        if not movimientos:
            QMessageBox.warning(self, "Sin movimientos", "No hay movimientos en este turno.")
            return
        import openpyxl
        from openpyxl.styles import Font
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"Turno_{self.id_apertura_actual}"
        ws['A1'] = "Reporte de turno"
        ws['A3'] = "Fecha/Hora"
        ws['B3'] = "Tipo"
        ws['C3'] = "Descripción"
        ws['D3'] = "Monto"
        ws['E3'] = "Usuario"
        for i, m in enumerate(movimientos, start=4):
            ws.cell(row=i, column=1, value=str(m['fecha_hora']))
            ws.cell(row=i, column=2, value=m['tipo_movimiento'])
            ws.cell(row=i, column=3, value=m['descripcion'])
            ws.cell(row=i, column=4, value=float(m['monto']))
            ws.cell(row=i, column=5, value=m['usuario_nombre'])
        ruta, _ = QFileDialog.getSaveFileName(self, "Guardar turno", f"Turno_{self.id_apertura_actual}.xlsx", "Excel files (*.xlsx)")
        if ruta:
            wb.save(ruta)
            QMessageBox.information(self, "Éxito", "Turno exportado")